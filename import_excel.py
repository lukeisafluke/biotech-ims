import os
import sys
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor

DB_HOST = os.environ.get("DB_HOST", "localhost")
DB_PORT = os.environ.get("DB_PORT", "5432")
DB_NAME = os.environ.get("DB_NAME", "inventory_db")
DB_USER = os.environ.get("DB_USER", "postgres")
DB_PASS = os.environ.get("DB_PASS", "postgres")

HEADER_CAT_SYNONYMS = {'cat. no.', 'cat#', 'sku', 'product no', 'product no.', 'catalog', 'catalog no', 'catalog no.'}


def get_db():
    return psycopg2.connect(
        dbname=DB_NAME,
        user=DB_USER,
        password=DB_PASS,
        host=DB_HOST,
        port=DB_PORT
    )


def parse_cat_no(cat_str):
    """
    Splits a catalog number into (sku, sub_sku).
    Product numbers with hyphens indicate SKU-subsku (e.g., 'R711-01' -> ('R711', '01')).
    Product numbers without hyphens have sub_sku = '' (e.g., 'EXS101' -> ('EXS101', '')).
    """
    cat_str = str(cat_str or '').strip()
    if not cat_str:
        return None, None
    if cat_str.lower() in HEADER_CAT_SYNONYMS:
        return None, None
    if '-' in cat_str:
        sku, sub_sku = cat_str.split('-', 1)
        return sku.strip(), sub_sku.strip()
    return cat_str, ''


def _find_col_index(headers, candidates):
    """Case-insensitive header lookup."""
    norm_headers = [str(h or '').strip().lower() for h in headers]
    for cand in candidates:
        cand_lower = cand.lower()
        for idx, h in enumerate(norm_headers):
            if h == cand_lower:
                return idx
    for cand in candidates:
        cand_lower = cand.lower()
        for idx, h in enumerate(norm_headers):
            if cand_lower in h:
                return idx
    return None


def import_excel_file(filepath, conn=None, supplier='Primary Supplier'):
    """
    Scans all worksheets in an Excel workbook for product catalog columns
    (Cat. No. / SKU, Product Name, Size, CAD price) and syncs with the database:
    - Inserts only new (sku, sub_sku) entries.
    - Updates existing (sku, sub_sku) entries ONLY when name, size, price, or supplier changed.
    - Skips unchanged products without duplicating or overwriting.
    """
    own_conn = False
    if conn is None:
        conn = get_db()
        own_conn = True

    supplier_clean = (supplier or 'Primary Supplier').strip() or 'Primary Supplier'
    wb = openpyxl.load_workbook(filepath, data_only=True)

    products_added = 0
    products_updated = 0
    products_unchanged = 0
    sheets_matched = []

    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        # Load existing products into a lookup map to avoid unnecessary writes
        cur.execute("SELECT product_id, sku, sub_sku, name, size, price, supplier FROM products")
        existing_map = {}
        for row in cur.fetchall():
            key = (row['sku'], row['sub_sku'] or '')
            existing_map[key] = {
                'product_id': row['product_id'],
                'name': row['name'] or '',
                'size': row['size'] or '',
                'price': round(float(row['price'] or 0.0), 2),
                'supplier': row['supplier'] or 'Primary Supplier'
            }

        seen_in_workbook = set()

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.max_row < 2:
                continue

            # Locate header row within the first 10 rows
            header_row_idx = None
            headers = []
            for r in range(1, min(11, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                cat_c = _find_col_index(row_vals, ['cat. no.', 'cat#', 'sku', 'product no', 'catalog'])
                cad_c = _find_col_index(row_vals, ['cad', 'price', 'unit price (cad)'])
                name_c = _find_col_index(row_vals, ['product name', 'name', 'item', 'description'])
                if cat_c is not None and (cad_c is not None or name_c is not None):
                    header_row_idx = r
                    headers = row_vals
                    break

            if not headers or header_row_idx is None:
                continue

            cat_idx = _find_col_index(headers, ['cat. no.', 'cat#', 'sku', 'product no', 'catalog'])
            name_idx = _find_col_index(headers, ['product name', 'name', 'item', 'description'])
            size_idx = _find_col_index(headers, ['size', 'unit', 'pack'])
            cad_idx = _find_col_index(headers, ['cad', 'price', 'unit price (cad)'])

            if cat_idx is None or cad_idx is None:
                continue

            sheets_matched.append(sheet_name)
            last_product_name = ''

            for r in range(header_row_idx + 1, ws.max_row + 1):
                raw_cat = ws.cell(r, cat_idx + 1).value
                if raw_cat is None or not str(raw_cat).strip():
                    continue

                sku, sub_sku = parse_cat_no(raw_cat)
                if not sku:
                    continue

                key = (sku, sub_sku)
                if key in seen_in_workbook:
                    continue
                seen_in_workbook.add(key)

                # Product Name (carry forward if merged/blank in grouped rows)
                raw_name = ws.cell(r, name_idx + 1).value if name_idx is not None else None
                if raw_name and str(raw_name).strip():
                    pname = str(raw_name).strip()
                    last_product_name = pname
                else:
                    pname = last_product_name or sku

                # Size
                raw_size = ws.cell(r, size_idx + 1).value if size_idx is not None else None
                size_str = str(raw_size).strip() if raw_size is not None and str(raw_size).strip() else 'Standard'

                # CAD Price
                raw_cad = ws.cell(r, cad_idx + 1).value if cad_idx is not None else 0.0
                if raw_cad is None or str(raw_cad).strip() == '' or str(raw_cad).strip().lower() == 'inquiry':
                    price_val = 0.0
                else:
                    try:
                        price_val = round(float(str(raw_cad).replace(',', '').replace('$', '').strip()), 2)
                    except (ValueError, TypeError):
                        price_val = 0.0

                existing = existing_map.get(key)
                if existing is None:
                    cur.execute("""
                        INSERT INTO products (sku, sub_sku, name, size, price, supplier)
                        VALUES (%s, %s, %s, %s, %s, %s)
                        RETURNING product_id
                    """, (sku, sub_sku, pname, size_str, price_val, supplier_clean))
                    new_id = cur.fetchone()['product_id']
                    existing_map[key] = {
                        'product_id': new_id,
                        'name': pname,
                        'size': size_str,
                        'price': price_val,
                        'supplier': supplier_clean
                    }
                    products_added += 1
                else:
                    name_changed = (existing['name'] != pname)
                    size_changed = (existing['size'] != size_str)
                    price_changed = (abs(existing['price'] - price_val) > 0.001)
                    supplier_changed = (existing['supplier'] != supplier_clean)

                    if name_changed or size_changed or price_changed or supplier_changed:
                        cur.execute("""
                            UPDATE products
                            SET name = %s,
                                size = %s,
                                price = %s,
                                supplier = %s
                            WHERE product_id = %s
                        """, (pname, size_str, price_val, supplier_clean, existing['product_id']))
                        existing.update({
                            'name': pname,
                            'size': size_str,
                            'price': price_val,
                            'supplier': supplier_clean
                        })
                        products_updated += 1
                    else:
                        products_unchanged += 1

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        if own_conn:
            conn.close()

    return {
        'products_added': products_added,
        'products_updated': products_updated,
        'products_unchanged': products_unchanged,
        'sheets_matched': sheets_matched,
        'total_processed': products_added + products_updated + products_unchanged
    }


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python import_excel.py <path_to_excel_file.xlsx> [supplier_name]")
        sys.exit(1)

    target_path = sys.argv[1]
    supplier_arg = sys.argv[2] if len(sys.argv) > 2 else 'Primary Supplier'
    if not os.path.exists(target_path):
        print(f"Error: Excel file not found at '{target_path}'")
        print("Usage: python import_excel.py <path_to_excel_file.xlsx> [supplier_name]")
        sys.exit(1)

    print(f"Scanning all sheets in '{target_path}' (Supplier: {supplier_arg}) ...")
    summary = import_excel_file(target_path, supplier=supplier_arg)
    print(
        f"Done! Matched sheets: {', '.join(summary['sheets_matched'])}\n"
        f"Total SKUs scanned: {summary['total_processed']} "
        f"({summary['products_added']} new added, "
        f"{summary['products_updated']} updated, "
        f"{summary['products_unchanged']} unchanged)."
    )
